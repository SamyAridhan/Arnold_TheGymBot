"""Offline stand-ins for Telegram and the sheet, used by tests and the demo."""
from __future__ import annotations

from .store import HEADERS, LOGGED, MemoryTable
from .telegram import TelegramError


class FakeTelegram:
    def __init__(self):
        self.messages: dict[int, dict] = {}
        self.next_id = 100
        self.toasts: list[str] = []
        self.deleted: list[int] = []

    def send(self, chat_id, text, kb=None):
        self.next_id += 1
        self.messages[self.next_id] = {"text": text, "kb": kb}
        return self.next_id

    def edit(self, chat_id, message_id, text, kb=None):
        if message_id not in self.messages:
            raise TelegramError("message to edit not found")
        if self.messages[message_id] == {"text": text, "kb": kb}:
            raise TelegramError("Bad Request: message is not modified")
        self.messages[message_id] = {"text": text, "kb": kb}

    def answer(self, query_id, text=""):
        self.toasts.append(text)

    def delete(self, chat_id, message_id):
        self.deleted.append(message_id)
        self.messages.pop(message_id, None)

    def set_commands(self, commands):
        pass

    def updates(self, offset, timeout):
        return []


def memory_tables(seed: dict[str, list[dict]] | None = None) -> tuple[dict, dict]:
    seed = seed or {}
    real = {k: MemoryTable(HEADERS[k], seed.get(k)) for k in HEADERS}
    test = {k: MemoryTable(HEADERS[k]) for k in LOGGED}
    return real, test


def load_xlsx(path: str) -> dict[str, list[dict]]:
    """Read the exported gym-log.xlsx so tests run against your real data."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = {}
    for key, title in (("sets", "Sets"), ("sessions", "Sessions"), ("notes", "Notes"),
                       ("program", "Program"), ("exercises", "Exercises")):
        ws = wb[title]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        out[key] = [{h: ("" if v is None else v) for h, v in zip(hdr, r)} for r in rows[1:]]
    return out
